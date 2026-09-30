import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import nsdecls, qn

def set_cell_margins(cell, top=100, bottom=100, left=150, right=150):
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = OxmlElement('w:tcMar')
    for m, val in [('top', top), ('bottom', bottom), ('left', left), ('right', right)]:
        node = OxmlElement(f'w:{m}')
        node.set(qn('w:w'), str(val))
        node.set(qn('w:type'), 'dxa')
        tcMar.append(node)
    tcPr.append(tcMar)

def set_cell_shading(cell, color_hex):
    shading_elm = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{color_hex}"/>')
    cell._tc.get_or_add_tcPr().append(shading_elm)

def set_table_borders(table):
    tblPr = table._tbl.tblPr
    borders = parse_xml(
        f'<w:tblBorders {nsdecls("w")}>'
        f'<w:top w:val="single" w:sz="6" w:space="0" w:color="CCCCCC"/>'
        f'<w:bottom w:val="single" w:sz="6" w:space="0" w:color="CCCCCC"/>'
        f'<w:left w:val="none"/>'
        f'<w:right w:val="none"/>'
        f'<w:insideH w:val="single" w:sz="4" w:space="0" w:color="E0E0E0"/>'
        f'<w:insideV w:val="none"/>'
        f'</w:tblBorders>'
    )
    tblPr.append(borders)

def build_lab03_doc(output_path):
    doc = docx.Document()

    # Page setup - Margins (1 inch = 72 pt)
    sections = doc.sections
    for section in sections:
        section.top_margin = Inches(0.8)
        section.bottom_margin = Inches(0.8)
        section.left_margin = Inches(0.8)
        section.right_margin = Inches(0.8)

    # Style configuration
    normal_style = doc.styles['Normal']
    normal_style.font.name = 'Times New Roman'
    normal_style.font.size = Pt(12)
    normal_style.font.color.rgb = RGBColor(0x22, 0x22, 0x22)
    normal_style.paragraph_format.line_spacing = 1.15
    normal_style.paragraph_format.space_after = Pt(4)

    # --- HEADER SECTION ---
    p_header = doc.add_paragraph()
    p_header.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r_uni = p_header.add_run("TRƯỜNG ĐẠI HỌC ĐÀ LẠT\n")
    r_uni.font.name = 'Times New Roman'
    r_uni.font.size = Pt(11)
    r_uni.font.bold = True
    r_uni.font.color.rgb = RGBColor(0x44, 0x44, 0x44)

    r_title = p_header.add_run("BÁO CÁO LAB 03 - MÔN PHÁT TRIỂN ỨNG DỤNG WEB NÂNG CAO\n")
    r_title.font.name = 'Times New Roman'
    r_title.font.size = Pt(15)
    r_title.font.bold = True
    r_title.font.color.rgb = RGBColor(0x11, 0x33, 0x66)

    r_sub = p_header.add_run("Đề tài nhóm: Culinary Blog – Blog Ẩm thực và Nấu ăn\n")
    r_sub.font.name = 'Times New Roman'
    r_sub.font.size = Pt(12)
    r_sub.font.italic = True

    # Info box paragraph
    p_info = doc.add_paragraph()
    p_info.paragraph_format.space_after = Pt(8)
    r_info = p_info.add_run(
        "Họ và tên: Nguyễn Phạm Phú Nam          MSSV: 2312695          Nhóm: 09\n"
        "Branch thực hiện: 2312695-NguyenPhamPhuNam-Lab3-Official\n"
        "Thời gian thực hiện: 23/09/2026 – 30/09/2026"
    )
    r_info.font.name = 'Times New Roman'
    r_info.font.size = Pt(11.5)
    r_info.font.bold = True

    p_note = doc.add_paragraph()
    p_note.paragraph_format.space_after = Pt(8)
    r_note = p_note.add_run(
        "Lưu ý: Báo cáo cá nhân thể hiện đầy đủ 4 yêu cầu tối thiểu của Lab 3 theo quy định LMS: "
        "(1) Domain Exceptions, (2) Repository & Unit of Work, (3) Ít nhất 2 API endpoints/thành viên, "
        "(4) Middleware xử lý ngoại lệ toàn cục trả về Problem Details (RFC 7807)."
    )
    r_note.font.name = 'Times New Roman'
    r_note.font.size = Pt(10.5)
    r_note.font.italic = True
    r_note.font.color.rgb = RGBColor(0x55, 0x55, 0x55)

    # --- TABLE OF ASSIGNED WORK ---
    p_tbl_title = doc.add_paragraph()
    r_tbl_title = p_tbl_title.add_run("I. BẢNG TỔNG HỢP CÔNG VIỆC THỰC HIỆN")
    r_tbl_title.font.name = 'Times New Roman'
    r_tbl_title.font.size = Pt(13)
    r_tbl_title.font.bold = True
    r_tbl_title.font.color.rgb = RGBColor(0x11, 0x33, 0x66)

    table_data = [
        (
            "1",
            "Cài đặt hệ thống Domain Exceptions\n"
            "• Xây dựng lớp ngoại lệ cơ sở DomainException tại tầng Domain (CulinaryBlog.Domain/Exceptions/DomainException.cs).\n"
            "• Xây dựng InvalidTokenStateException kế thừa DomainException mô hình hóa vi phạm quy tắc vòng đời của RefreshToken.\n"
            "• Đóng gói phương thức nghiệp vụ Revoke(DateTimeOffset, string?) trên Entity RefreshToken, tự động ném InvalidTokenStateException khi vi phạm kiểm tra trạng thái.\n"
            "• Tái sử dụng các Application Exceptions hiện hữu: NotFoundException, UnauthorizedException, ConflictException, AccountLockedException, InvalidGoogleTokenException.\n"
            "• Bổ sung Unit Tests (DomainExceptionTests.cs) kiểm chứng 100% logic ném ngoại lệ miền.",
            "backend/src/CulinaryBlog.Domain/Exceptions/DomainException.cs\n"
            "backend/src/CulinaryBlog.Domain/Exceptions/InvalidTokenStateException.cs\n"
            "backend/src/CulinaryBlog.Domain/Entities/RefreshToken.cs\n"
            "backend/tests/CulinaryBlog.Application.Tests/Domain/DomainExceptionTests.cs",
            "100%"
        ),
        (
            "2",
            "Cài đặt Repository Pattern cho Refresh Token\n"
            "• Duy trì tính trừu tượng IRefreshTokenRepository đặt tại tầng Application, bảo đảm tầng nghiệp vụ không phụ thuộc trực tiếp vào ORM/EF Core.\n"
            "• Cài đặt RefreshTokenRepository tại tầng Infrastructure với đầy đủ phương thức async và CancellationToken: SaveRefreshTokenAsync, GetByHashAsync, UpdateRefreshTokenAsync, TryRotateRefreshTokenAsync.\n"
            "• Bảo toàn cơ chế xoay vòng Refresh Token nguyên tử (Atomic Refresh Token Rotation) chống race condition double-refresh trên PostgreSQL (qua ExecuteUpdateAsync & Database Transaction) và đồng bộ an toàn trên InMemory Test Provider.",
            "backend/src/CulinaryBlog.Application/Contracts/Authentication/IRefreshTokenRepository.cs\n"
            "backend/src/CulinaryBlog.Infrastructure/Repositories/RefreshTokenRepository.cs\n"
            "backend/tests/CulinaryBlog.Integration.Tests/Endpoints/AuthEndpointsTests.cs",
            "100%"
        ),
        (
            "3",
            "Cài đặt Unit of Work Pattern\n"
            "• Xây dựng interface trừu tượng IUnitOfWork tại tầng Application (CulinaryBlog.Application/Contracts/Persistence/IUnitOfWork.cs).\n"
            "• Cài đặt UnitOfWork tại tầng Infrastructure (CulinaryBlog.Infrastructure/Persistence/UnitOfWork.cs) bao bọc AuthDbContext, cung cấp SaveChangesAsync(CancellationToken).\n"
            "• Đăng ký IUnitOfWork trong Dependency Injection (AddScoped) tại Infrastructure/DependencyInjection.cs.\n"
            "• Tích hợp thực tế IUnitOfWork vào LogoutCommandHandler để đồng bộ và lưu trữ trạng thái thu hồi token.\n"
            "• Viết kiểm thử đơn vị và kiểm thử tích hợp (UnitOfWorkTests.cs, UnitOfWorkIntegrationTests.cs) xác thực lưu dữ liệu thành công.",
            "backend/src/CulinaryBlog.Application/Contracts/Persistence/IUnitOfWork.cs\n"
            "backend/src/CulinaryBlog.Infrastructure/Persistence/UnitOfWork.cs\n"
            "backend/src/CulinaryBlog.Infrastructure/DependencyInjection.cs\n"
            "backend/src/CulinaryBlog.Application/Features/Auth/Logout/LogoutCommandHandler.cs\n"
            "backend/tests/CulinaryBlog.Application.Tests/Persistence/UnitOfWorkTests.cs\n"
            "backend/tests/CulinaryBlog.Integration.Tests/Persistence/UnitOfWorkIntegrationTests.cs",
            "100%"
        ),
        (
            "4",
            "Cài đặt các API Endpoints Tuần 3 (3 endpoints / thành viên)\n"
            "• POST /api/v1/auth/refresh (FR-AUTH-004): Cấp phát cặp Access Token mới (15m) và xoay vòng Refresh Token (7d), phát hiện tái sử dụng token (reuse detection), thu hồi token con trực tiếp, đảm bảo tính nguyên tử chống double refresh (1 thành công 200, 1 từ chối 401).\n"
            "• POST /api/v1/auth/logout (FR-AUTH-005): Đăng xuất người dùng, thu hồi Refresh Token, bảo đảm tính idempotent (204) và bảo vệ token của tài khoản khác.\n"
            "• GET /api/v1/auth/me (FR-AUTH-006): Lấy hồ sơ người dùng đang xác thực với DTO an toàn, lọc bỏ hoàn toàn các trường nhạy cảm của ASP.NET Identity (PasswordHash, SecurityStamp).",
            "backend/src/CulinaryBlog.API/Endpoints/AuthEndpoints.cs\n"
            "backend/src/CulinaryBlog.Application/Features/Auth/TokenRefresh/*\n"
            "backend/src/CulinaryBlog.Application/Features/Auth/Logout/*\n"
            "backend/src/CulinaryBlog.Application/Features/Auth/Profile/*\n"
            "backend/src/CulinaryBlog.Application/DTOs/Auth/*",
            "100%"
        ),
        (
            "5",
            "Xây dựng Global Exception Middleware & Problem Details (RFC 7807)\n"
            "• Tạo GlobalExceptionMiddleware tại tầng API (CulinaryBlog.API/Middleware/GlobalExceptionMiddleware.cs) bắt toàn bộ ngoại lệ chưa được xử lý ở mức toàn cục.\n"
            "• Chuẩn hóa phản hồi theo tiêu chuẩn RFC 7807 Problem Details (Content-Type: application/problem+json) với đầy đủ trường: type, title, status, detail, instance, traceId.\n"
            "• Ánh xạ mã lỗi: ValidationException -> 400 (kèm danh sách lỗi từng trường), NotFoundException -> 404, UnauthorizedException -> 401, ConflictException -> 409, AccountLockedException -> 423, DomainException -> 400, Unexpected Exception -> 500 (che giấu stack trace và thông tin nhạy cảm).\n"
            "• Đăng ký middleware ở đầu pipeline của Program.cs; tinh gọn AuthEndpoints.cs để ngoại lệ lan truyền tự nhiên tới middleware.\n"
            "• Bổ sung Integration Tests (GlobalExceptionMiddlewareTests.cs) kiểm thử toàn diện các luồng lỗi.",
            "backend/src/CulinaryBlog.API/Middleware/GlobalExceptionMiddleware.cs\n"
            "backend/src/CulinaryBlog.API/Program.cs\n"
            "backend/src/CulinaryBlog.API/Endpoints/AuthEndpoints.cs\n"
            "backend/tests/CulinaryBlog.Integration.Tests/Endpoints/GlobalExceptionMiddlewareTests.cs",
            "100%"
        ),
        (
            "6",
            "Giao diện Frontend /profile và Kiểm thử E2E Playwright\n"
            "• Xây dựng màn hình xem hồ sơ /profile bằng Next.js 15, Tailwind CSS, TypeScript.\n"
            "• Hỗ trợ cơ chế tự động xoay vòng refresh token khi nhận 401 Unauthorized và tự động thử lại yêu cầu (refresh-and-retry).\n"
            "• Chức năng Đăng xuất: xóa sạch sessionStorage và điều hướng về /auth/login.\n"
            "• Cô lập helper API trong frontend/src/lib/auth-api.ts, không gây xung đột api.ts dùng chung.\n"
            "• Viết kịch bản kiểm thử E2E (auth-profile.spec.ts), chạy Playwright thành công 6/6 kịch bản.",
            "frontend/src/app/profile/page.tsx\n"
            "frontend/src/lib/auth-api.ts\n"
            "frontend/e2e/auth-profile.spec.ts",
            "100%"
        ),
        (
            "7",
            "Kiểm thử hệ thống toàn diện (Backend & Frontend)\n"
            "• Backend Tests: 97/97 tests PASS (Architecture: 3/3, Application: 69/69, Integration: 25/25).\n"
            "• Frontend Lint: PASS (0 lỗi, 0 cảnh báo ESLint).\n"
            "• Frontend Build: PASS (biên dịch Next.js hoàn tất thành công).\n"
            "• Playwright E2E: PASS (6/6 kịch bản trình duyệt thành công).",
            "backend/tests/*\n"
            "frontend/*",
            "100%"
        )
    ]

    # Create table
    table = doc.add_table(rows=1, cols=4)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    set_table_borders(table)

    headers = ["STT", "Nội dung công việc thực hiện", "File / Minh chứng kỹ thuật", "Tiến độ %"]
    col_widths = [Inches(0.5), Inches(3.4), Inches(2.3), Inches(0.7)]

    # Format header row
    hdr_cells = table.rows[0].cells
    for i, title in enumerate(headers):
        hdr_cells[i].text = title
        hdr_cells[i].width = col_widths[i]
        set_cell_margins(hdr_cells[i], top=120, bottom=120, left=100, right=100)
        set_cell_shading(hdr_cells[i], "E8EEF5")
        p = hdr_cells[i].paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        for run in p.runs:
            run.font.name = 'Times New Roman'
            run.font.size = Pt(11)
            run.font.bold = True
            run.font.color.rgb = RGBColor(0x11, 0x33, 0x66)

    # Populate rows
    for row_idx, data in enumerate(table_data):
        row = table.add_row()
        cells = row.cells
        for col_idx in range(4):
            cells[col_idx].width = col_widths[col_idx]
            cells[col_idx].text = data[col_idx]
            set_cell_margins(cells[col_idx], top=100, bottom=100, left=100, right=100)
            if row_idx % 2 == 1:
                set_cell_shading(cells[col_idx], "F9FAFC")

            p = cells[col_idx].paragraphs[0]
            if col_idx in (0, 3):
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            else:
                p.alignment = WD_ALIGN_PARAGRAPH.LEFT

            for run in p.runs:
                run.font.name = 'Times New Roman'
                run.font.size = Pt(10.5)

    doc.add_paragraph().paragraph_format.space_after = Pt(8)

    # --- SECTION II: CHI TIẾT CÔNG VIỆC ĐÃ HOÀN THÀNH ---
    p_sec2 = doc.add_paragraph()
    r_sec2 = p_sec2.add_run("II. CHI TIẾT CÁC NỘI DUNG ĐÃ HOÀN THÀNH")
    r_sec2.font.name = 'Times New Roman'
    r_sec2.font.size = Pt(13)
    r_sec2.font.bold = True
    r_sec2.font.color.rgb = RGBColor(0x11, 0x33, 0x66)

    details = [
        ("1. Hệ thống Domain Exceptions và Application Exceptions",
         "Trong kiến trúc Clean Architecture, việc phân định rõ trách nhiệm xử lý lỗi giữa tầng Domain và tầng Application là rất quan trọng:\n"
         "• Tầng Domain: Em đã thiết lập lớp ngoại lệ trừu tượng cơ sở DomainException tại thư mục CulinaryBlog.Domain/Exceptions. Đồng thời, em cài đặt lớp cụ thể InvalidTokenStateException để đại diện cho các vi phạm quy tắc bất biến trong vòng đời của thực thể RefreshToken. Trên entity RefreshToken, em bổ sung phương thức nghiệp vụ Revoke(DateTimeOffset, string?) giúp đóng gói trạng thái; nếu một token đã bị thu hồi trước đó nhưng tiếp tục nhận yêu cầu thu hồi, phương thức sẽ chủ động ném InvalidTokenStateException. Điều này đảm bảo tính toàn vẹn dữ liệu ngay tại trung tâm miền nghiệp vụ.\n"
         "• Tầng Application: Tiếp tục duy trì các ngoại lệ phục vụ luồng điều phối ca sử dụng (Use Case) tại CulinaryBlog.Application/Common/Exceptions bao gồm: NotFoundException (thực thể không tồn tại trong hệ thống), UnauthorizedException (thông tin xác thực không hợp lệ, token hết hạn hoặc phát hiện tái sử dụng token), ConflictException (xung đột dữ liệu như email đã đăng ký), AccountLockedException (tài khoản bị khóa sau 5 lần đăng nhập thất bại) và InvalidGoogleTokenException.\n"
         "Cách phân chia này giúp mã nguồn tuân thủ nguyên tắc Dependency Rule: tầng Domain hoàn toàn độc lập, không phụ thuộc vào ASP.NET Core hay thư viện bên ngoài."),

        ("2. Mô hình Repository và tính nguyên tử của Refresh Token Rotation",
         "Em tiếp tục áp dụng và hoàn thiện mô hình Repository Pattern cho thực thể RefreshToken:\n"
         "• Interface IRefreshTokenRepository được đặt tại tầng Application (CulinaryBlog.Application/Contracts/Authentication/IRefreshTokenRepository.cs), đóng vai trò là bản đặc tả hợp đồng truy xuất dữ liệu mà các Use Case (Command/Query Handler) dựa vào. Tầng Application hoàn toàn không biết chi tiết công nghệ lưu trữ bên dưới.\n"
         "• Lớp cài đặt RefreshTokenRepository được đặt tại tầng Infrastructure (CulinaryBlog.Infrastructure/Repositories/RefreshTokenRepository.cs), sử dụng EF Core và DbContext để thao tác cơ sở dữ liệu PostgreSQL.\n"
         "• Điểm nhấn kỹ thuật quan trọng: Để giải quyết triệt để lỗi chạy đồng thời (race condition double-refresh) khi hai yêu cầu refresh cùng gửi đồng thời một refresh token hợp lệ, em đã cài đặt phương thức TryRotateRefreshTokenAsync. Trên PostgreSQL, phương thức này thực hiện một câu lệnh UPDATE có điều kiện (WHERE Id = @id AND RevokedAt IS NULL AND ReplacedByTokenHash IS NULL) và kiểm tra số dòng bị ảnh hưởng (affected rows). Chỉ duy nhất một request chuyển trạng thái token thành công (affected == 1) mới được phép thêm token con mới trong cùng một transaction; request còn lại nhận kết quả thất bại (affected == 0), transaction rollback và hệ thống từ chối với HTTP 401 Unauthorized."),

        ("3. Mô hình Unit of Work (UoW)",
         "Để đáp ứng chuẩn mực thiết kế và tiêu chí bắt buộc của Lab 3, em đã bổ sung mẫu thiết kế Unit of Work vào hệ thống:\n"
         "• Abstraction: Định nghĩa interface IUnitOfWork tại tầng Application (CulinaryBlog.Application/Contracts/Persistence/IUnitOfWork.cs) cung cấp phương thức bất đồng bộ SaveChangesAsync(CancellationToken).\n"
         "• Implementation: Cài đặt lớp UnitOfWork tại tầng Infrastructure (CulinaryBlog.Infrastructure/Persistence/UnitOfWork.cs) bao bọc AuthDbContext, đóng vai trò là điểm chốt chặn ghi nhận các thay đổi dữ liệu tập trung.\n"
         "• Sử dụng thực tế: Em đã đăng ký IUnitOfWork vào DI container và tiêm (inject) vào LogoutCommandHandler. Khi người dùng thực hiện đăng xuất hợp lệ, handler gọi phương thức Revoke() trên entity, chuyển qua Repository để đánh dấu và gọi UnitOfWork.SaveChangesAsync() để đảm bảo tính nhất quán của giao dịch.\n"
         "• Kiểm thử: Đã xây dựng các bài kiểm thử đơn vị (UnitOfWorkTests.cs) và kiểm thử tích hợp (UnitOfWorkIntegrationTests.cs) xác nhận việc lưu trữ dữ liệu vào DbContext diễn ra chính xác."),

        ("4. Cài đặt các API Endpoints thuộc phân hệ Xác thực (Auth)",
         "Em đã hoàn thành đầy đủ 3 API Endpoints quan trọng phục vụ quản lý phiên người dùng:\n"
         "• POST /api/v1/auth/refresh: Tiếp nhận Refresh Token thô, băm SHA-256 để so khớp trong cơ sở dữ liệu. Kiểm tra hạn dùng, kiểm tra trạng thái hoạt động của tài khoản người dùng, thực hiện xoay vòng token nguyên tử, cấp cặp Access Token mới (15 phút) và Refresh Token mới (7 ngày). Khi phát hiện token đã bị thu hồi được sử dụng lại (Reuse Detection), hệ thống sẽ lập tức thu hồi token con trực tiếp để vô hiệu hóa phiên của kẻ tấn công.\n"
         "• POST /api/v1/auth/logout: Nhận Refresh Token và xác thực danh tính người dùng qua JWT Bearer claim. Đảm bảo tính idempotent (gọi nhiều lần vẫn trả 204 NoContent) và ngăn chặn người dùng đăng xuất nhầm token của tài khoản khác.\n"
         "• GET /api/v1/auth/me: Yêu cầu JWT hợp lệ, truy vấn thông tin tài khoản và trả về UserProfileDto an toàn, phục vụ hiển thị trên giao diện người dùng."),

        ("5. Global Exception Middleware và chuẩn Problem Details (RFC 7807)",
         "Thay vì bắt lỗi phân tán tại từng Endpoint bằng try-catch thủ công, em đã xây dựng GlobalExceptionMiddleware tại tầng API (CulinaryBlog.API/Middleware/GlobalExceptionMiddleware.cs):\n"
         "• Middleware này bao bọc toàn bộ pipeline xử lý request: try { await _next(context); } catch (Exception ex) { ... }.\n"
         "• Định dạng phản hồi: Tất cả lỗi trả về đều tuân thủ chuẩn RFC 7807 với Content-Type: application/problem+json, bao gồm các thuộc tính tiêu chuẩn: type, title, status, detail, instance và traceId.\n"
         "• Bảng ánh xạ ngoại lệ: ValidationException -> 400 Bad Request kèm danh sách lỗi chi tiết; NotFoundException -> 404 Not Found; UnauthorizedException -> 401 Unauthorized; ConflictException -> 409 Conflict; AccountLockedException -> 423 Locked; DomainException -> 400 Bad Request.\n"
         "• An toàn thông tin: Đối với các ngoại lệ hệ thống không lường trước (Unhandled Exceptions), middleware ghi log lỗi chi tiết trên máy chủ nhưng trả về HTTP 500 với thông điệp chung an toàn: 'An unexpected error occurred while processing your request.', tuyệt đối không để lộ stack trace, chuỗi kết nối hay chi tiết cơ sở dữ liệu cho client.\n"
         "• Đã cấu hình middleware tại Program.cs và tinh gọn toàn bộ AuthEndpoints.cs để các typed exceptions lan truyền tự nhiên tới middleware xử lý."),

        ("6. Giao diện người dùng /profile và Kiểm thử E2E",
         "• Xây dựng màn hình /profile bằng Next.js 15, TypeScript và Tailwind CSS, hiển thị các trường thông tin thực tế: Email, Tên hiển thị (DisplayName), Vai trò (Roles), Trạng thái xác thực email (EmailConfirmed), Ngày tham gia (CreatedAt), Ảnh đại diện (AvatarUrl) và Tiểu sử (Bio).\n"
         "• Tích hợp logic tự động bắt lỗi 401: khi access token hết hạn, client tự động gọi API refresh để lấy token mới rồi tự động tải lại dữ liệu mà người dùng không bị gián đoạn trải nghiệm.\n"
         "• Tách riêng các hàm gọi API sang frontend/src/lib/auth-api.ts để tránh xung đột với mã nguồn của các thành viên khác trong nhóm.\n"
         "• Kiểm thử tự động E2E với Playwright (auth-profile.spec.ts) bao quát các kịch bản: chuyển hướng khi chưa đăng nhập, hiển thị thông tin thành công, tự động refresh khi token hết hạn và đăng xuất xóa phiên làm việc.")
    ]

    for title, content in details:
        p_t = doc.add_paragraph()
        r_t = p_t.add_run(title)
        r_t.font.name = 'Times New Roman'
        r_t.font.size = Pt(12)
        r_t.font.bold = True
        r_t.font.color.rgb = RGBColor(0x11, 0x33, 0x66)

        p_c = doc.add_paragraph()
        p_c.paragraph_format.line_spacing = 1.15
        p_c.paragraph_format.space_after = Pt(6)
        r_c = p_c.add_run(content)
        r_c.font.name = 'Times New Roman'
        r_c.font.size = Pt(11)

    # --- SECTION III: NỘI DUNG CHƯA HOÀN THÀNH & HẠN CHẾ ---
    p_sec3 = doc.add_paragraph()
    r_sec3 = p_sec3.add_run("III. NỘI DUNG CHƯA HOÀN THÀNH VÀ HẠN CHẾ KỸ THUẬT")
    r_sec3.font.name = 'Times New Roman'
    r_sec3.font.size = Pt(13)
    r_sec3.font.bold = True
    r_sec3.font.color.rgb = RGBColor(0x11, 0x33, 0x66)

    p_limit = doc.add_paragraph()
    p_limit.paragraph_format.line_spacing = 1.15
    p_limit.paragraph_format.space_after = Pt(6)
    r_limit = p_limit.add_run(
        "1. Về yêu cầu tối thiểu của Lab 3: Em đã hoàn thành đầy đủ 100% cả 4 tiêu chí bắt buộc theo quy định của môn học (Domain Exceptions, Repository, Unit of Work, 3 API endpoints, Global Exception Middleware trả về Problem Details).\n\n"
        "2. Một số hạn chế kỹ thuật thực tế và nội dung theo kế hoạch tuần tiếp theo:\n"
        "• Cơ chế thu hồi họ token khi phát hiện tái sử dụng (Reuse Detection): Hiện tại hệ thống phát hiện việc sử dụng lại token cũ và thu hồi được token con trực tiếp (direct descendant - 1 hop). Cơ chế thu hồi đệ quy đa cấp toàn bộ nhánh cây token (full recursive token-family revocation) chưa được áp dụng trong tuần này nhằm tránh rủi ro suy giảm hiệu năng và độ phức tạp truy vấn.\n"
        "• Mã trạng thái HTTP cho lỗi Validation: Hiện tại hệ thống đang trả về mã HTTP 400 Bad Request (kèm Problem Details) theo hợp đồng giao tiếp tạm thời của nhóm; mâu thuẫn giữa 400 và 422 trong tài liệu SRS (CONFLICT-011) sẽ được nhóm thống nhất xử lý trong các giai đoạn sau.\n"
        "• Chức năng chỉnh sửa thông tin cá nhân (FR-AUTH-007): Theo bảng phân công công việc của nhóm, tính năng chỉnh sửa hồ sơ người dùng (PATCH /api/v1/auth/me) thuộc phạm vi của tuần sau nên chưa triển khai trong Lab 3."
    )
    r_limit.font.name = 'Times New Roman'
    r_limit.font.size = Pt(11)

    # --- SECTION IV: KẾT QUẢ KIỂM THỬ THỰC TẾ ---
    p_sec4 = doc.add_paragraph()
    r_sec4 = p_sec4.add_run("IV. KẾT QUẢ KIỂM THỬ THỰC TẾ (TEST RESULTS)")
    r_sec4.font.name = 'Times New Roman'
    r_sec4.font.size = Pt(13)
    r_sec4.font.bold = True
    r_sec4.font.color.rgb = RGBColor(0x11, 0x33, 0x66)

    p_test = doc.add_paragraph()
    p_test.paragraph_format.line_spacing = 1.15
    p_test.paragraph_format.space_after = Pt(6)
    r_test = p_test.add_run(
        "Toàn bộ các bộ kiểm thử tự động đã được thực thi và xác minh trực tiếp trên máy phát triển:\n"
        "• Backend Architecture Tests: 3 / 3 test PASS (Đảm bảo cấu trúc Clean Architecture và tính toàn vẹn phụ thuộc giữa các tầng).\n"
        "• Backend Application Tests: 69 / 69 test PASS (Bao gồm các bài test Command, Validator, Handler của Đăng ký, Đăng nhập, Google Login, Refresh Token, Logout, Profile, Domain Exception và Unit of Work).\n"
        "• Backend Integration Tests: 25 / 25 test PASS (Bao gồm kiểm thử API thực tế, rate limiting, đăng nhập, xoay vòng refresh token, thu hồi token, kiểm tra race condition chạy đồng thời và kiểm thử GlobalExceptionMiddleware Problem Details).\n"
        "• Tổng số Backend Tests: 97 / 97 test PASS (100% đạt, không có test nào thất bại hoặc bị bỏ qua).\n"
        "• Frontend ESLint: PASS (Không phát hiện bất kỳ lỗi cú pháp hoặc cảnh báo linter nào).\n"
        "• Frontend Build: PASS (Next.js 15 biên dịch production thành công, xuất bản đầy đủ các trang tĩnh và động).\n"
        "• Playwright E2E Tests: 6 / 6 test PASS (Bao gồm kiểm thử quy trình Đăng ký, Đăng nhập, Xem hồ sơ, Tự động xoay vòng refresh token khi nhận 401 và Đăng xuất)."
    )
    r_test.font.name = 'Times New Roman'
    r_test.font.size = Pt(11)

    # --- SECTION V: TỰ ĐÁNH GIÁ ---
    p_sec5 = doc.add_paragraph()
    r_sec5 = p_sec5.add_run("V. TỰ ĐÁNH GIÁ MỨC ĐỘ HOÀN THÀNH")
    r_sec5.font.name = 'Times New Roman'
    r_sec5.font.size = Pt(13)
    r_sec5.font.bold = True
    r_sec5.font.color.rgb = RGBColor(0x11, 0x33, 0x66)

    p_eval = doc.add_paragraph()
    p_eval.paragraph_format.line_spacing = 1.15
    p_eval.paragraph_format.space_after = Pt(12)
    r_eval = p_eval.add_run(
        "Căn cứ vào 4 yêu cầu tối thiểu chính thức của Lab 3 (Domain Exceptions, Repository & Unit of Work, ít nhất 2 API endpoints, Global Exception Middleware Problem Details) và kết quả kiểm thử toàn diện thực tế:\n\n"
        "Tự đánh giá mức độ hoàn thành Lab 3: 100%"
    )
    r_eval.font.name = 'Times New Roman'
    r_eval.font.size = Pt(12)
    r_eval.font.bold = True
    r_eval.font.color.rgb = RGBColor(0x00, 0x55, 0x00)

    # Save document
    doc.save(output_path)
    print("Document successfully created.")

if __name__ == '__main__':
    out_file = r"D:\Nam 4\Phát triển ứng dụng Web nâng cao\Nhom9_Culinary_Blog\BaoCao_Lab3\Lab03_2312695_NguyenPhamPhuNam.docx"
    build_lab03_doc(out_file)
